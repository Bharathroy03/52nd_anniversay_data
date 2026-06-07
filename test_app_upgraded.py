import os
import unittest
from unittest.mock import patch, MagicMock
import json
import io
from openpyxl import load_workbook

# Set dummy environment variables before importing Config and App
os.environ['SUPABASE_URL'] = 'https://abc.supabase.co'
os.environ['SUPABASE_KEY'] = 'dummy-key-99999'
os.environ['SECRET_KEY'] = 'test-session-key'

from config import Config
from app import app

class UpgradedAppTestCase(unittest.TestCase):
    def setUp(self):
        # Create Flask test client
        self.app = app.test_client()
        self.app.testing = True

    def test_directory_setup(self):
        """Test that the application automatically creates the required folder structure."""
        dirs = [
            'static/css',
            'static/js',
            'static/images',
            'templates',
            'database'
        ]
        for d in dirs:
            self.assertTrue(os.path.exists(d), f"Folder {d} was not automatically created.")

    def test_super_admin_login_success(self):
        """Test Super Admin (Bharath Kumar) login sets correct session role."""
        credentials = {
            "username": "bharath",
            "password": "Admin@123"
        }
        with self.app as client:
            response = client.post('/api/admin/login', 
                                   data=json.dumps(credentials),
                                   content_type='application/json')
            self.assertEqual(response.status_code, 200)
            data = json.loads(response.data)
            self.assertTrue(data['success'])
            
            with client.session_transaction() as sess:
                self.assertEqual(sess.get('logged_in'), True)
                self.assertEqual(sess.get('role'), 'super_admin')

    def test_admin_store_head_login_success(self):
        """Test Admin & Store Head (Saddam Husain) login sets correct session role."""
        credentials = {
            "username": "saddam",
            "password": "Saddam@123"
        }
        with self.app as client:
            response = client.post('/api/admin/login', 
                                   data=json.dumps(credentials),
                                   content_type='application/json')
            self.assertEqual(response.status_code, 200)
            data = json.loads(response.data)
            self.assertTrue(data['success'])
            
            with client.session_transaction() as sess:
                self.assertEqual(sess.get('logged_in'), True)
                self.assertEqual(sess.get('role'), 'admin_store_head')

    def test_admin_api_protection(self):
        """Test that admin endpoints restrict unauthorized session access."""
        res = self.app.get('/api/admin/dashboard')
        self.assertEqual(res.status_code, 401)

        res2 = self.app.get('/api/admin/customers')
        self.assertEqual(res2.status_code, 401)

        res3 = self.app.get('/api/admin/export-excel')
        self.assertEqual(res3.status_code, 401)

        res4 = self.app.get('/api/admin/settings')
        self.assertEqual(res4.status_code, 401)

    @patch('app.get_supabase')
    def test_rbac_settings_access(self, mock_get_supabase):
        """Test settings page Role-Based Access Control restrictions."""
        # 1. Super Admin access (expect 200)
        with self.app as client:
            with client.session_transaction() as sess:
                sess['logged_in'] = True
                sess['role'] = 'super_admin'
            
            res = client.get('/api/admin/settings')
            self.assertEqual(res.status_code, 200)
            data = json.loads(res.data)
            self.assertTrue(data['success'])
            self.assertIn("Super Admin system configuration settings", data['message'])

        # 2. Admin & Store Head access (expect 403)
        with self.app as client:
            with client.session_transaction() as sess:
                sess['logged_in'] = True
                sess['role'] = 'admin_store_head'
            
            res = client.get('/api/admin/settings')
            self.assertEqual(res.status_code, 403)
            data = json.loads(res.data)
            self.assertFalse(data['success'])
            self.assertIn("Access Denied", data['message'])

    @patch('app.get_supabase')
    def test_admin_dashboard_with_filters(self, mock_get_supabase):
        """Test dashboard stats endpoint fetches and filters metrics."""
        with self.app as client:
            with client.session_transaction() as sess:
                sess['logged_in'] = True
                sess['role'] = 'super_admin'
                
            mock_db = MagicMock()
            mock_get_supabase.return_value = mock_db
            
            # Chain-friendly mock query builder
            mock_query = MagicMock()
            mock_db.table.return_value.select.return_value = mock_query
            mock_query.eq.return_value = mock_query
            mock_query.gte.return_value = mock_query
            mock_query.lte.return_value = mock_query
            mock_query.or_.return_value = mock_query
            mock_query.order.return_value = mock_query
            
            mock_query.execute.return_value.data = [
                {"id": 1, "app_registration_status": "Completed", "invitation_issue": True},
                {"id": 2, "app_registration_status": "Pending", "invitation_issue": False},
                {"id": 3, "app_registration_status": "Not Interested", "invitation_issue": False}
            ]
            
            # Request with filters
            res = client.get('/api/admin/dashboard?store_id=1&status=Completed&date=2026-06-07')
            self.assertEqual(res.status_code, 200)
            result = json.loads(res.data)
            self.assertTrue(result['success'])
            self.assertEqual(result['stats']['total'], 3)
            self.assertEqual(result['stats']['completed'], 1)
            self.assertEqual(result['stats']['pending'], 1)
            self.assertEqual(result['stats']['not_interested'], 1)
            self.assertEqual(result['stats']['issues'], 1)
            self.assertEqual(result['stats']['branches'], 13)
            self.assertEqual(result['user']['role'], 'super_admin')

    @patch('app.get_supabase')
    def test_admin_export_excel(self, mock_get_supabase):
        """Test exporting formatted Excel workbook with 3 worksheets."""
        with self.app as client:
            with client.session_transaction() as sess:
                sess['logged_in'] = True
                sess['role'] = 'super_admin'
                
            mock_db = MagicMock()
            mock_get_supabase.return_value = mock_db
            
            # Chain-friendly mocks for both customer_entries and stores queries
            mock_query = MagicMock()
            mock_query.select.return_value = mock_query
            mock_query.eq.return_value = mock_query
            mock_query.gte.return_value = mock_query
            mock_query.lte.return_value = mock_query
            mock_query.or_.return_value = mock_query
            mock_query.order.return_value = mock_query
            
            mock_stores_query = MagicMock()
            mock_stores_query.select.return_value = mock_stores_query
            mock_stores_query.order.return_value = mock_stores_query
            
            def table_side_effect(table_name):
                if table_name == 'customer_entries':
                    return mock_query
                elif table_name == 'stores':
                    return mock_stores_query
                return mock_query
                
            mock_db.table.side_effect = table_side_effect
            
            # Setup database mock values
            mock_entries = [
                {
                    "id": 1,
                    "store_id": 2,
                    "customer_name": "John Doe",
                    "mobile_number": "9876543210",
                    "app_registration_status": "Completed",
                    "invitation_issue": True,
                    "issue_description": "Barcode issue",
                    "created_at": "2026-06-07T10:00:00Z",
                    "stores": {"id": 2, "store_name": "Nelamangala - 1 - (Bus Stand)"}
                },
                {
                    "id": 2,
                    "store_id": 3,
                    "customer_name": "Jane Doe",
                    "mobile_number": "0123456789",
                    "app_registration_status": "Pending",
                    "invitation_issue": False,
                    "issue_description": None,
                    "created_at": "2026-06-07T11:00:00Z",
                    "stores": {"id": 3, "store_name": "Chintamani - 1 - (GSLN Theatre)"}
                }
            ]
            
            mock_stores = [
                {"id": 2, "store_name": "Nelamangala - 1 - (Bus Stand)"},
                {"id": 3, "store_name": "Chintamani - 1 - (GSLN Theatre)"}
            ]
            
            # Setup mock execute values
            mock_query.execute.return_value.data = mock_entries
            mock_stores_query.execute.return_value.data = mock_stores
            
            res = client.get('/api/admin/export-excel?date=2026-06-07')
            self.assertEqual(res.status_code, 200)
            self.assertEqual(res.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            self.assertIn("attachment; filename=team_saddam_zone_data_", res.headers.get('Content-Disposition', ''))
            
            # Load and verify Excel sheet contents
            excel_bytes = io.BytesIO(res.data)
            wb = load_workbook(excel_bytes)
            
            self.assertIn("All Submissions", wb.sheetnames)
            self.assertIn("Store Wise Summary", wb.sheetnames)
            self.assertIn("Registration Status Summary", wb.sheetnames)
            
            # Verify sheet 1 details
            ws1 = wb["All Submissions"]
            self.assertEqual(ws1.cell(row=1, column=1).value, "Submission ID")
            self.assertEqual(ws1.cell(row=1, column=4).value, "Mobile Number")
            
            # Verify text formatted mobile number (stored as string to preserve leading zero)
            self.assertIsInstance(ws1.cell(row=2, column=4).value, str)
            self.assertEqual(ws1.cell(row=2, column=4).value, "0123456789")
            
            # Verify sheet 2 details
            ws2 = wb["Store Wise Summary"]
            self.assertEqual(ws2.cell(row=1, column=1).value, "Store Name")
            self.assertEqual(ws2.cell(row=1, column=2).value, "Total Submissions")
            
            # Verify sheet 3 details
            ws3 = wb["Registration Status Summary"]
            self.assertEqual(ws3.cell(row=1, column=1).value, "App Registration Status")
            self.assertEqual(ws3.cell(row=5, column=1).value, "Total")

    @patch('app.get_supabase')
    def test_api_admin_stats(self, mock_get_supabase):
        """Test retrieving stats metrics for the dashboard."""
        mock_db = MagicMock()
        mock_get_supabase.return_value = mock_db
        
        mock_stores = [{"id": 1, "store_name": "Store A"}]
        mock_entries = [
            {
                "id": 1,
                "store_id": 1,
                "customer_name": "Alice",
                "mobile_number": "1234567890",
                "app_registration_status": "Completed",
                "invitation_issue": False,
                "issue_description": None,
                "created_at": "2026-06-07T12:00:00Z"
            }
        ]
        
        # Mocking the table queries for stores and customer_entries
        mock_stores_query = MagicMock()
        mock_stores_query.select.return_value.order.return_value.execute.return_value.data = mock_stores
        
        mock_entries_query = MagicMock()
        mock_entries_query.select.return_value.execute.return_value.data = mock_entries
        
        def table_side_effect(name):
            if name == 'stores':
                return mock_stores_query
            elif name == 'customer_entries':
                return mock_entries_query
            return MagicMock()
            
        mock_db.table.side_effect = table_side_effect
        
        with self.app as client:
            with client.session_transaction() as sess:
                sess['logged_in'] = True
                sess['role'] = 'super_admin'
                
            res = client.get('/api/admin/stats')
            self.assertEqual(res.status_code, 200)
            data = json.loads(res.data)
            self.assertTrue(data['success'])
            self.assertEqual(data['total_customers'], 1)
            self.assertEqual(data['todays_customers'], 1)
            self.assertEqual(len(data['day_wise_customers']), 1)
            self.assertEqual(len(data['store_wise_customers']), 1)

    @patch('app.get_supabase')
    def test_api_update_customer(self, mock_get_supabase):
        """Test updating a customer entry with self-healing fallback for updated_at."""
        mock_db = MagicMock()
        mock_get_supabase.return_value = mock_db
        
        mock_update_query = MagicMock()
        # First call fails because of missing updated_at column, second succeeds
        mock_update_query.eq.return_value.execute.side_effect = [
            Exception("column \"updated_at\" of relation \"customer_entries\" does not exist"),
            MagicMock(data=[{"id": 1}])
        ]
        
        mock_db.table.return_value.update.return_value = mock_update_query
        
        payload = {
            "store_id": 1,
            "customer_name": "Bob Updated",
            "mobile_number": "9998887776",
            "app_registration_status": "Pending",
            "invitation_issue": "No",
            "issue_description": ""
        }
        
        with self.app as client:
            with client.session_transaction() as sess:
                sess['logged_in'] = True
                sess['role'] = 'admin_store_head'
                
            res = client.put('/api/admin/customer/1', data=json.dumps(payload), content_type='application/json')
            self.assertEqual(res.status_code, 200)
            data = json.loads(res.data)
            self.assertTrue(data['success'])
            
            # Verify update was called twice due to retry fallback mechanism
            update_calls = mock_db.table.return_value.update.call_args_list
            self.assertEqual(len(update_calls), 2)
            self.assertIn("updated_at", update_calls[0][0][0])
            self.assertNotIn("updated_at", update_calls[1][0][0])

    @patch('app.get_supabase')
    def test_api_delete_customer(self, mock_get_supabase):
        """Test deleting an individual customer entry."""
        mock_db = MagicMock()
        mock_get_supabase.return_value = mock_db
        
        mock_delete_query = MagicMock()
        mock_delete_query.eq.return_value.execute.return_value.data = [{"id": 5}]
        mock_db.table.return_value.delete.return_value = mock_delete_query
        
        with self.app as client:
            with client.session_transaction() as sess:
                sess['logged_in'] = True
                sess['role'] = 'admin_store_head'
                
            res = client.delete('/api/admin/customer/5')
            self.assertEqual(res.status_code, 200)
            data = json.loads(res.data)
            self.assertTrue(data['success'])
            mock_delete_query.eq.assert_called_with('id', 5)

    @patch('app.get_supabase')
    def test_api_delete_all_customers_rbac(self, mock_get_supabase):
        """Test deleting all customer records restrictions."""
        mock_db = MagicMock()
        mock_get_supabase.return_value = mock_db
        
        # 1. Store head should get forbidden response
        with self.app as client:
            with client.session_transaction() as sess:
                sess['logged_in'] = True
                sess['role'] = 'admin_store_head'
                
            res = client.delete('/api/admin/customers/delete-all')
            self.assertEqual(res.status_code, 403)
            
        # 2. Super admin is authorized
        mock_delete_query = MagicMock()
        mock_delete_query.neq.return_value.execute.return_value.data = [{"id": 10}]
        mock_db.table.return_value.delete.return_value = mock_delete_query
        
        with self.app as client:
            with client.session_transaction() as sess:
                sess['logged_in'] = True
                sess['role'] = 'super_admin'
                
            res = client.delete('/api/admin/customers/delete-all')
            self.assertEqual(res.status_code, 200)
            data = json.loads(res.data)
            self.assertTrue(data['success'])

if __name__ == '__main__':
    unittest.main()
